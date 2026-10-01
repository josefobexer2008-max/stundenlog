"""Liest Projektzeitplan.xlsx (Vertex42-Gantt-Vorlage) und schreibt js/seed.js.

Aufruf:  python3 tools/excel_to_seed.py [pfad/zu/Projektzeitplan.xlsx]
Spalte B = Name, D = Fortschritt, E/F = Start/Ende (nur bei Phasen).
Die Prozentwerte aus der Excel werden der Gruppe "Obexer & Taferner" (g2) zugeordnet.
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

# Offensichtliche Tippfehler aus der Excel
FIX = {
    "Technikraum Positionieren (UG)": "Technikraum positionieren (UG)",
    "Erschliesung Keller-EG festlegen": "Erschließung Keller–EG festlegen",
    "Erschliesungskonzept festlegen (EG)": "Erschließungskonzept festlegen (EG)",
    "Gruppensitzung mit Verbesserungsvorschläge": "Gruppensitzung mit Verbesserungsvorschlägen",
    "Vorenwurf abschliesen durch gegenseitige Kontrolle": "Vorentwurf abschließen durch gegenseitige Kontrolle",
    "Raumgrößen Kontrolieren": "Raumgrößen kontrollieren",
    "Möbelierung einzeichnen": "Möblierung einzeichnen",
    "Verkehrsflächen Optimieren": "Verkehrsflächen optimieren",
    "Erschließung Optimieren": "Erschließung optimieren",
    "Spannweiten Prüffen": "Spannweiten prüfen",
    "Fensteranordnung optimieren": "Fensteranordnung optimieren",
    "Fassaden- und Dachgestallung": "Fassaden- und Dachgestaltung",
    "Wirtschaftlichkeit prüffen": "Wirtschaftlichkeit prüfen",
    "Geländemodel modellieren": "Geländemodell modellieren",
    "OG2/DG modelieren": "OG2/DG modellieren",
    "Qualitätskontrolle des BIM models durch Schulinternene Person": "Qualitätskontrolle des BIM-Modells durch schulinterne Person",
    "Einreichplanung abschliesen": "Einreichplanung abschließen",
    "Detailzeichungen erstellen": "Detailzeichnungen erstellen",
    "Nachweis für Luftdichtheit erbingen": "Nachweis für Luftdichtheit erbringen",
    "PV- Anlage": "PV-Anlage",
    "Sicherheitspplan erstellen": "Sicherheitsplan erstellen",
    "Vorstellen vor Ausgewähler Kommision (mit Fachlehrern)": "Vorstellen vor ausgewählter Kommission (mit Fachlehrern)",
    "Freigabe zug abgabe des Projekts": "Freigabe zur Abgabe des Projekts",
    "3D Modell für Renderings vorbereiten": "3D-Modell für Renderings vorbereiten",
}
# Falsche Daten in der Excel (Phase -> (Start, Ende))
DATE_FIX = {"Detailplanung": (None, dt.date(2027, 2, 7))}
PLACEHOLDER = re.compile(r"^Aufgabe \d+$")


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Projektplan"]
    phases, subs, progress = [], [], []
    cur = None
    for row in range(8, ws.max_row + 1):
        name = ws.cell(row, 2).value
        if not isinstance(name, str) or not name.strip():
            continue
        name = re.sub(r"\s+", " ", name).strip()
        if name.startswith("Neue Zeilen"):
            break
        start, end = ws.cell(row, 5).value, ws.cell(row, 6).value
        if isinstance(start, dt.datetime) and isinstance(end, dt.datetime):
            s, e = start.date(), end.date()
            fs, fe = DATE_FIX.get(name, (None, None))
            s, e = fs or s, fe or e
            cur = {"id": f"p{len(phases) + 1}", "name": name, "start_date": s.isoformat(),
                   "end_date": e.isoformat(), "description": "", "sort": len(phases) + 1}
            phases.append(cur)
            continue
        if cur is None or PLACEHOLDER.match(name):
            continue
        n = sum(1 for x in subs if x["phase_id"] == cur["id"]) + 1
        sid = f"{cur['id']}s{n:02d}"
        subs.append({"id": sid, "phase_id": cur["id"], "name": FIX.get(name, name), "description": "", "sort": n})
        pct = ws.cell(row, 4).value
        if isinstance(pct, (int, float)) and pct > 0:
            progress.append({"id": f"{sid}:g2", "subphase_id": sid, "group_id": "g2",
                             "percent": int(round(pct * 100 / 5) * 5)})
    groups = [
        {"id": "g1", "name": "Engl & Bürgstaller", "members": ["Matthias Engl", "Paul Bürgstaller"], "sort": 1},
        {"id": "g2", "name": "Obexer & Taferner", "members": ["Josef Obexer", "Felix Taferner"], "sort": 2},
    ]
    seed = {"groups": groups, "phases": phases, "subphases": subs, "progress": progress}
    out = Path(__file__).resolve().parent.parent / "js" / "seed.js"
    out.write_text("// Erzeugt mit tools/excel_to_seed.py aus Projektzeitplan.xlsx – Startdaten beim ersten Öffnen.\n"
                   "window.SEED = " + json.dumps(seed, ensure_ascii=False, indent=1) + ";\n", encoding="utf-8")
    print(f"{len(phases)} Phasen, {len(subs)} Unterphasen, {len(progress)} Fortschrittswerte -> {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name("Projektzeitplan.xlsx")))
